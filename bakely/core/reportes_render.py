"""
BAKELY · Motores de renderizado de informes (independientes del ORM)

Cada función recibe datos planos (dict/list) y devuelve bytes, así se pueden
probar sin base de datos y reutilizar desde vistas, comandos o tareas programadas.

  pdf_ebr(datos)            → Informe Oficial de Lote (eBR) con membrete UPEC
  pdf_manifiesto(datos)     → Manifiesto de carga y despacho al CEDI
  pdf_kardex(datos)         → Kardex y rendimiento WMS
  xlsx_auditoria(filas)     → Bitácora de auditoría de operarios
  xlsx_kardex(datos)        → Kardex + rendimiento + IRA (varias hojas)
  csv_auditoria(filas)      → Bitácora en CSV (UTF-8 con BOM para Excel)
"""
import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.graphics.barcode.qr import QrCodeWidget
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

# Paleta institucional
AZUL = colors.HexColor("#0F172A")
VERDE_UPEC = colors.HexColor("#0E7C4A")
AMBAR = colors.HexColor("#D97706")
ROJO = colors.HexColor("#B91C1C")
GRIS = colors.HexColor("#64748B")
GRIS_CLARO = colors.HexColor("#F1F5F9")

_ss = getSampleStyleSheet()
ST = {
    "h1": ParagraphStyle("h1", parent=_ss["Heading1"], fontSize=15, textColor=AZUL, spaceAfter=2, leading=18),
    "h2": ParagraphStyle("h2", parent=_ss["Heading2"], fontSize=11, textColor=VERDE_UPEC, spaceBefore=10,
                         spaceAfter=4),
    "p": ParagraphStyle("p", parent=_ss["Normal"], fontSize=8.5, leading=11),
    "small": ParagraphStyle("s", parent=_ss["Normal"], fontSize=7, leading=9, textColor=GRIS),
    "cell": ParagraphStyle("c", parent=_ss["Normal"], fontSize=7.5, leading=9),
    "center": ParagraphStyle("ce", parent=_ss["Normal"], fontSize=8, alignment=TA_CENTER),
    "th": ParagraphStyle("th", parent=_ss["Normal"], fontSize=7.5, leading=9, textColor=colors.white,
                         fontName="Helvetica-Bold"),
}


# =============================================================================
# Utilidades PDF
# =============================================================================
def _qr(texto, size=22 * mm):
    w = QrCodeWidget(texto)
    b = w.getBounds()
    d = Drawing(size, size, transform=[size / (b[2] - b[0]), 0, 0, size / (b[3] - b[1]), 0, 0])
    d.add(w)
    return d


def _membrete(titulo, subtitulo, institucion, planta):
    def dibujar(canvas, doc):
        canvas.saveState()
        w, h = doc.pagesize
        # Banda superior
        canvas.setFillColor(AZUL)
        canvas.rect(0, h - 22 * mm, w, 22 * mm, fill=1, stroke=0)
        canvas.setFillColor(VERDE_UPEC)
        canvas.rect(0, h - 23.5 * mm, w, 1.5 * mm, fill=1, stroke=0)
        # Sello UPEC (círculo con siglas)
        canvas.setStrokeColor(colors.white)
        canvas.setFillColor(VERDE_UPEC)
        canvas.circle(18 * mm, h - 11 * mm, 7.5 * mm, fill=1, stroke=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 8)
        canvas.drawCentredString(18 * mm, h - 12.2 * mm, "UPEC")
        canvas.setFont("Helvetica-Bold", 12)
        canvas.drawString(29 * mm, h - 9.5 * mm, institucion.upper())
        canvas.setFont("Helvetica", 8.5)
        canvas.drawString(29 * mm, h - 14.5 * mm, f"{planta} · Sistema BAKELY WMS/MES")
        canvas.setFont("Helvetica-Bold", 9)
        canvas.drawRightString(w - 12 * mm, h - 9.5 * mm, titulo)
        canvas.setFont("Helvetica", 7.5)
        canvas.drawRightString(w - 12 * mm, h - 14.5 * mm, subtitulo)
        # Pie
        canvas.setFillColor(GRIS)
        canvas.setFont("Helvetica", 6.5)
        canvas.drawString(12 * mm, 8 * mm,
                          f"Documento generado electrónicamente por BAKELY · {datetime.now():%Y-%m-%d %H:%M} · "
                          "Las firmas electrónicas se registran en la bitácora de auditoría.")
        canvas.drawRightString(w - 12 * mm, 8 * mm, f"Página {doc.page}")
        canvas.restoreState()
    return dibujar


def _tabla(data, col_widths=None, header=True, zebra=True, estilos_extra=None):
    filas = [[c if hasattr(c, "wrap") else Paragraph(str(c if c is not None else "—"), ST["cell"]) for c in fila]
             for fila in data]
    if header:
        filas[0] = [Paragraph(str(h), ST["th"]) for h in data[0]]
    t = Table(filas, colWidths=col_widths, repeatRows=1 if header else 0)
    estilo = [
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if header:
        estilo += [("BACKGROUND", (0, 0), (-1, 0), AZUL)]
    if zebra:
        for r in range(1 if header else 0, len(data)):
            if r % 2 == 0:
                estilo.append(("BACKGROUND", (0, r), (-1, r), GRIS_CLARO))
    t.setStyle(TableStyle(estilo + (estilos_extra or [])))
    return t


def _kv(pares, col=2):
    """Tabla clave/valor en `col` pares por fila."""
    filas, fila = [], []
    for k, v in pares:
        fila += [Paragraph(f"<font color='#64748B'>{k}</font>", ST["cell"]),
                 Paragraph(f"<b>{v if v not in (None, '') else '—'}</b>", ST["cell"])]
        if len(fila) == col * 2:
            filas.append(fila)
            fila = []
    if fila:
        filas.append(fila + [""] * (col * 2 - len(fila)))
    t = Table(filas)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
                           ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    return t


def _sello(conforme):
    if conforme is True:
        return "<font color='#0E7C4A'><b>CONFORME</b></font>"
    if conforme is False:
        return "<font color='#B91C1C'><b>NO CONFORME</b></font>"
    return "<font color='#D97706'><b>PENDIENTE</b></font>"


def _firmas(firmantes):
    celdas = []
    for f in firmantes:
        celdas.append(Paragraph(
            f"<br/><br/>______________________________<br/><b>{f.get('nombre') or '(sin firmar)'}</b><br/>"
            f"<font color='#64748B'>{f['cargo']}</font><br/><font size='6' color='#64748B'>{f.get('fecha') or ''}</font>",
            ST["center"]))
    t = Table([celdas])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return KeepTogether([Spacer(1, 6 * mm), t])


def _doc(buf, titulo, subtitulo, meta, pagesize=A4):
    doc = SimpleDocTemplate(buf, pagesize=pagesize, topMargin=30 * mm, bottomMargin=16 * mm,
                            leftMargin=12 * mm, rightMargin=12 * mm, title=titulo, author="BAKELY · UPEC")
    cb = _membrete(titulo, subtitulo, meta.get("institucion", "Universidad Politécnica Estatal del Carchi"),
                   meta.get("planta", "Planta de Almidones y Licores"))
    return doc, cb


# =============================================================================
# 1. INFORME OFICIAL DE LOTE (eBR)
# =============================================================================
def pdf_ebr(d: dict) -> bytes:
    """
    d = {
      meta:{institucion, planta, url_trazabilidad},
      lote:{codigo, producto, sku, linea, estado, cantidad, unidad, elaboracion, caducidad, abv, humedad,
            cenizas, pcc1_conforme, certificado, liberado_por, fecha_liberacion},
      orden:{codigo, cantidad_planificada, unidad, inicio, cierre, equipo},
      origenes:[{lote, producto, proveedor, procedencia, cantidad}],
      registros:[{fecha, estacion, etapa, control, variables, conforme, desviaciones, registrado_por, aprobado_por}],
      cortes:[{alambique, carga_l, abv_carga, cabezas_l, corazon_l, abv_corazon, colas_l, metanol, conforme, motivos}],
      iot_resumen:[{equipo, variable, min, max, promedio, n, alertas}],
      balance:{...}, firmas:[{cargo, nombre, fecha}]
    }
    """
    buf = io.BytesIO()
    lote = d["lote"]
    doc, cb = _doc(buf, "INFORME OFICIAL DE LOTE · eBR", f"Lote {lote['codigo']}", d.get("meta", {}))
    el = []

    cab = Table([[
        [Paragraph(f"{lote['producto']}", ST["h1"]),
         Paragraph(f"SKU {lote['sku']} · {lote['linea']} · Estado: <b>{lote['estado']}</b>", ST["p"]),
         Paragraph(f"Resultado PCC/PC global: {_sello(d.get('conforme_global'))}", ST["p"])],
        _qr(d.get("meta", {}).get("url_trazabilidad", lote["codigo"])),
    ]], colWidths=[150 * mm, 36 * mm])
    cab.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    el.append(cab)

    el.append(Paragraph("1. Identificación del lote", ST["h2"]))
    el.append(_kv([
        ("Código de lote", lote["codigo"]), ("Cantidad", f"{lote['cantidad']} {lote['unidad']}"),
        ("Fecha de elaboración", lote["elaboracion"]), ("Fecha de caducidad", lote.get("caducidad")),
        ("Orden de producción", d.get("orden", {}).get("codigo")), ("Equipo principal", d.get("orden", {}).get("equipo")),
        ("Graduación final (% ABV)", lote.get("abv")), ("Humedad final (%)", lote.get("humedad")),
        ("Cenizas (%)", lote.get("cenizas")), ("Certificado de inocuidad", lote.get("certificado")),
    ]))

    if d.get("origenes"):
        el.append(Paragraph("2. Genealogía · materias primas de origen", ST["h2"]))
        el.append(_tabla([["Lote MP", "Materia prima", "Productor", "Procedencia", "Cantidad"]] +
                         [[o["lote"], o["producto"], o.get("proveedor"), o.get("procedencia"), o.get("cantidad")]
                          for o in d["origenes"]], col_widths=[30 * mm, 48 * mm, 40 * mm, 38 * mm, 30 * mm]))

    el.append(Paragraph("3. Registro de etapas y puntos de control", ST["h2"]))
    filas = [["Fecha", "Estación / Etapa", "Control", "Variables", "Resultado", "Operario / Aprobó"]]
    for r in d.get("registros", []):
        vars_ = "<br/>".join(f"{k}: <b>{v}</b>" for k, v in (r.get("variables") or {}).items()) or "—"
        res = _sello(r["conforme"]) + ("<br/><font size='6' color='#B91C1C'>" + "; ".join(r["desviaciones"]) + "</font>"
                                       if r.get("desviaciones") else "")
        filas.append([r["fecha"], f"<b>{r['etapa']}</b><br/><font color='#64748B'>{r['estacion']}</font>",
                      r["control"], vars_, res, f"{r['registrado_por']}<br/><font color='#64748B'>{r.get('aprobado_por') or ''}</font>"])
    el.append(_tabla(filas, col_widths=[24 * mm, 38 * mm, 16 * mm, 46 * mm, 34 * mm, 28 * mm]))

    if d.get("cortes"):
        el.append(Paragraph("4. PCC-1 · Cortes de destilación (cabezas / corazón / colas)", ST["h2"]))
        filas = [["Alambique", "Carga (L @ %ABV)", "Cabezas (L)", "Corazón (L @ %ABV)", "Colas (L)",
                  "Metanol mg/100 mL a.a.", "Resultado"]]
        for c in d["cortes"]:
            filas.append([c["alambique"], f"{c['carga_l']} @ {c['abv_carga']}", c["cabezas_l"],
                          f"<b>{c['corazon_l']} @ {c['abv_corazon']}</b>", c["colas_l"], c.get("metanol"),
                          _sello(c["conforme"]) + (f"<br/><font size='6'>{'; '.join(c.get('motivos') or [])}</font>"
                                                   if c.get("motivos") else "")])
        el.append(_tabla(filas, col_widths=[22 * mm, 28 * mm, 20 * mm, 30 * mm, 18 * mm, 28 * mm, 40 * mm]))

    if d.get("iot_resumen"):
        el.append(Paragraph("5. Monitoreo IoT continuo (resumen estadístico)", ST["h2"]))
        el.append(_tabla([["Equipo", "Variable", "Mín", "Máx", "Promedio", "Lecturas", "Alertas"]] +
                         [[i["equipo"], i["variable"], i["min"], i["max"], i["promedio"], i["n"],
                           f"<font color='{'#B91C1C' if i['alertas'] else '#0E7C4A'}'><b>{i['alertas']}</b></font>"]
                          for i in d["iot_resumen"]]))

    if d.get("balance"):
        el.append(Paragraph("6. Balance de materia", ST["h2"]))
        etiquetas = {
            "entrada": "Entrada (consumo)", "salida": "Salida (producto)", "merma": "Mermas registradas",
            "no_contabilizado": "No contabilizado", "rendimiento_pct": "Rendimiento (%)",
            "lap_carga": "LAP cargado (L)", "lap_corazon": "LAP en corazón (L)", "lap_fracciones": "LAP total fracciones (L)",
            "recuperacion_corazon_pct": "Recuperación en corazón (%)", "perdida_lap_pct": "Pérdida de alcohol (%)",
        }
        el.append(_kv([(etiquetas.get(k, k), v) for k, v in d["balance"].items() if v is not None]))
        el.append(Paragraph("LAP = litros de alcohol puro = volumen × %ABV / 100.", ST["small"]))

    el.append(Paragraph("7. Liberación y firmas de aprobación", ST["h2"]))
    el.append(_kv([("Liberado por", lote.get("liberado_por")), ("Fecha de liberación", lote.get("fecha_liberacion"))]))
    el.append(_firmas(d.get("firmas") or [
        {"cargo": "Operario responsable de proceso"}, {"cargo": "Control de Calidad"},
        {"cargo": "Director de Planta UPEC"}]))
    doc.build(el, onFirstPage=cb, onLaterPages=cb)
    return buf.getvalue()


# =============================================================================
# 2. MANIFIESTO DE CARGA AL CEDI
# =============================================================================
def pdf_manifiesto(contrato: dict, meta: dict = None) -> bytes:
    buf = io.BytesIO()
    desp = contrato["despacho"]
    doc, cb = _doc(buf, "MANIFIESTO DE CARGA · CEDI", f"Despacho {desp['codigo']}", meta or {})
    el = [Paragraph(f"Despacho {desp['codigo']} hacia {contrato['destino']['cedi_codigo']}", ST["h1"]),
          Paragraph(f"Contrato <b>{contrato['tipo_mensaje']}</b> v{contrato['version_contrato']} · emitido "
                    f"{contrato['emitido_en'][:19].replace('T', ' ')}", ST["p"])]
    t = desp.get("transporte", {})
    el.append(Paragraph("Datos de transporte", ST["h2"]))
    el.append(_kv([("Fecha programada", desp["fecha_programada"][:16].replace("T", " ")),
                   ("Guía de remisión", desp.get("guia_remision")), ("Transportista", t.get("transportista")),
                   ("Placa", t.get("placa")), ("Conductor", t.get("conductor")),
                   ("Condición de temperatura", t.get("condicion_temperatura"))]))
    el.append(Paragraph("Detalle de carga", ST["h2"]))
    filas = [["SKU / Descripción", "Lote", "Cant.", "Emp.", "Palet", "Elab. / Cad.", "Calidad"]]
    for i in contrato["items"]:
        c = i["calidad"]
        cal = (f"{c.get('graduacion_alcoholica_abv')}% ABV · PCC-1 {'OK' if c.get('pcc1_conforme') else 'NO'}"
               if i["linea"] == "BEBIDA" else f"Humedad {c.get('humedad_pct')}%")
        filas.append([f"<b>{i['sku']}</b><br/>{i['descripcion']}", i["lote"], f"{i['cantidad']} {i['unidad']}",
                      i["unidades_empaque"], i.get("codigo_palet"),
                      f"{i['fecha_elaboracion']}<br/>{i.get('fecha_caducidad') or '—'}", cal])
    el.append(_tabla(filas, col_widths=[46 * mm, 26 * mm, 22 * mm, 12 * mm, 24 * mm, 22 * mm, 34 * mm]))
    tot = contrato["totales"]
    el.append(Spacer(1, 3 * mm))
    el.append(_kv([("Líneas", tot["lineas"]), ("Kg de almidón", tot["kg_almidon"]), ("Botellas", tot["botellas"])],
                  col=3))
    el.append(_firmas([{"cargo": "Despachó (Almacén Planta)"}, {"cargo": "Transportista"},
                       {"cargo": "Recibió (CEDI)"}]))
    doc.build(el, onFirstPage=cb, onLaterPages=cb)
    return buf.getvalue()


# =============================================================================
# 3. KARDEX (PDF)
# =============================================================================
def pdf_kardex(d: dict) -> bytes:
    """d = {meta, periodo, movimientos:[...], rendimiento:[...], ira, mermas:[{motivo, cantidad, unidad}]}"""
    buf = io.BytesIO()
    doc, cb = _doc(buf, "KARDEX Y RENDIMIENTO WMS", d.get("periodo", ""), d.get("meta", {}), pagesize=landscape(A4))
    el = [Paragraph("Kardex valorado por lote y rendimiento de planta", ST["h1"])]
    ira = d.get("ira")
    el.append(_kv([("Periodo", d.get("periodo")), ("Movimientos", len(d.get("movimientos", []))),
                   ("Exactitud de inventario (IRA)", f"{ira} %" if ira is not None else "Sin conteos"),
                   ("Meta IRA", ">= 95 %")], col=4))
    if d.get("rendimiento"):
        el.append(Paragraph("Rendimiento por orden de producción", ST["h2"]))
        el.append(_tabla([["Orden", "Producto", "Entrada", "Salida", "Merma", "Rendimiento %", "Recup. alcohol %"]] +
                         [[r["orden"], r["producto"], r["entrada"], r["salida"], r["merma"], r.get("rendimiento_pct"),
                           r.get("recuperacion_corazon_pct")] for r in d["rendimiento"]]))
    if d.get("mermas"):
        el.append(Paragraph("Mermas por motivo", ST["h2"]))
        el.append(_tabla([["Motivo", "Cantidad", "Unidad"]] + [[m["motivo"], m["cantidad"], m["unidad"]] for m in d["mermas"]]))
    el.append(PageBreak())
    el.append(Paragraph("Movimientos", ST["h2"]))
    el.append(_tabla([["Fecha", "Tipo", "SKU", "Lote", "Cantidad", "Origen", "Destino", "Saldo lote", "Usuario", "Ref."]] +
                     [[m["fecha"], m["tipo"], m["sku"], m["lote"], f"{m['signo']}{m['cantidad']} {m['unidad']}",
                       m.get("origen"), m.get("destino"), m["saldo"], m["usuario"], m.get("referencia")]
                      for m in d.get("movimientos", [])]))
    doc.build(el, onFirstPage=cb, onLaterPages=cb)
    return buf.getvalue()


# =============================================================================
# Excel / CSV
# =============================================================================
_H_FILL = PatternFill("solid", fgColor="0F172A")
_H_FONT = Font(bold=True, color="FFFFFF")
_BORDE = Border(bottom=Side(style="thin", color="CBD5E1"))


def _hoja(ws, titulo, encabezados, filas, anchos=None):
    ws.title = titulo[:31]
    ws.append(encabezados)
    for c in ws[1]:
        c.fill, c.font = _H_FILL, _H_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for f in filas:
        ws.append(list(f))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(encabezados, 1):
        ancho = (anchos or {}).get(h)
        if not ancho:
            muestra = [len(str(h))] + [len(str(f[i - 1])) for f in filas[:200] if i - 1 < len(f)]
            ancho = min(max(muestra) + 2, 60)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    return ws


CAMPOS_AUDITORIA = ["Fecha/hora", "Usuario", "Rol", "Estación usuario", "Acción", "Entidad", "ID", "Objeto",
                    "Estación del dato", "Cambios (antes → después)", "Sobrescritura admin", "IP", "Detalle"]


def _fila_auditoria(r):
    cambios = "; ".join(f"{k}: {v[0]} → {v[1]}" for k, v in (r.get("cambios") or {}).items())
    return [r["timestamp"], r["username"], r.get("rol"), r.get("estacion_usuario"), r["accion"], r.get("modelo"),
            r.get("objeto_id"), r.get("objeto_repr"), r.get("estacion_objeto"), cambios,
            "SÍ" if r.get("sobrescritura_admin") else "", r.get("ip"), r.get("detalle")]


def xlsx_auditoria(filas: list) -> bytes:
    wb = Workbook()
    ws = _hoja(wb.active, "Bitácora", CAMPOS_AUDITORIA, [_fila_auditoria(r) for r in filas],
               anchos={"Cambios (antes → después)": 60, "Objeto": 34, "Detalle": 40})
    n = len(filas) + 1
    # Resalta sobrescrituras administrativas e intentos denegados
    ws.conditional_formatting.add(f"A2:M{n}", FormulaRule(formula=['$K2="SÍ"'],
                                  fill=PatternFill("solid", fgColor="FEF3C7"), font=Font(bold=True, color="92400E")))
    ws.conditional_formatting.add(f"A2:M{n}", FormulaRule(formula=['$E2="DENEGADO"'],
                                  fill=PatternFill("solid", fgColor="FEE2E2")))
    # Hoja resumen con fórmulas vivas
    rs = wb.create_sheet("Resumen")
    rs.append(["Indicador", "Valor"])
    for c in rs[1]:
        c.fill, c.font = _H_FILL, _H_FONT
    rs.append(["Eventos totales", f"=COUNTA(Bitácora!A2:A{n})"])
    rs.append(["Sobrescrituras administrativas", f'=COUNTIF(Bitácora!K2:K{n},"SÍ")'])
    rs.append(["Intentos denegados", f'=COUNTIF(Bitácora!E2:E{n},"DENEGADO")'])
    rs.append(["Inicios de sesión fallidos", f'=COUNTIF(Bitácora!E2:E{n},"LOGIN_KO")'])
    rs.column_dimensions["A"].width, rs.column_dimensions["B"].width = 34, 14
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def csv_auditoria(filas: list) -> bytes:
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(CAMPOS_AUDITORIA)
    for r in filas:
        w.writerow(_fila_auditoria(r))
    return ("﻿" + out.getvalue()).encode("utf-8")


def xlsx_kardex(d: dict) -> bytes:
    wb = Workbook()
    movs = d.get("movimientos", [])
    _hoja(wb.active, "Kardex",
          ["Fecha", "Tipo", "SKU", "Producto", "Lote", "Entrada", "Salida", "Unidad", "Origen", "Destino",
           "Saldo lote", "Motivo merma", "Usuario", "Estación", "Referencia"],
          [[m["fecha"], m["tipo"], m["sku"], m.get("producto"), m["lote"],
            m["cantidad"] if m["signo"] == "+" else None, m["cantidad"] if m["signo"] == "-" else None,
            m["unidad"], m.get("origen"), m.get("destino"), m["saldo"], m.get("motivo_merma"), m["usuario"],
            m.get("estacion"), m.get("referencia")] for m in movs])
    n = len(movs) + 1
    ws = wb["Kardex"]
    ws.append([])
    ws.append(["TOTALES", "", "", "", "", f"=SUM(F2:F{n})", f"=SUM(G2:G{n})"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)

    _hoja(wb.create_sheet(), "Rendimiento",
          ["Orden", "Producto", "Línea", "Entrada", "Salida", "Merma", "Rendimiento %", "LAP carga", "LAP corazón",
           "Recuperación alcohol %"],
          [[r["orden"], r["producto"], r.get("linea"), r["entrada"], r["salida"], r["merma"], None,
            r.get("lap_carga"), r.get("lap_corazon"), None] for r in d.get("rendimiento", [])])
    rw = wb["Rendimiento"]
    for i in range(2, rw.max_row + 1):  # fórmulas vivas en lugar de valores fijos
        rw[f"G{i}"] = f'=IF(D{i}>0,ROUND(E{i}/D{i}*100,2),"")'
        rw[f"J{i}"] = f'=IF(N(H{i})>0,ROUND(I{i}/H{i}*100,2),"")'

    _hoja(wb.create_sheet(), "Existencias",
          ["Zona", "Ubicación", "SKU", "Lote", "Cantidad", "Unidad", "Caducidad", "Días para caducar", "Estado"],
          [[e["zona"], e["ubicacion"], e["sku"], e["lote"], e["cantidad"], e["unidad"], e.get("caducidad"),
            e.get("dias"), e.get("estado")] for e in d.get("existencias", [])])

    conteos = d.get("conteos", [])
    _hoja(wb.create_sheet(), "IRA",
          ["Fecha", "Ubicación", "Lote", "Sistema", "Físico", "Diferencia", "Dif. %", "Exacto (±2%)"],
          [[c["fecha"], c["ubicacion"], c["lote"], c["sistema"], c["fisico"], None, None, None] for c in conteos])
    iw = wb["IRA"]
    for i in range(2, iw.max_row + 1):
        iw[f"F{i}"] = f"=E{i}-D{i}"
        iw[f"G{i}"] = f'=IF(D{i}=0,IF(E{i}=0,0,100),ROUND(ABS(F{i})/D{i}*100,2))'
        iw[f"H{i}"] = f'=IF(G{i}<=2,"SÍ","NO")'
    m = iw.max_row
    iw.append([])
    iw.append(["IRA %", "", "", "", "", "", "", f'=IF(COUNTA(H2:H{m})=0,"",ROUND(COUNTIF(H2:H{m},"SÍ")/COUNTA(H2:H{m})*100,2))'])
    iw[f"A{iw.max_row}"].font = Font(bold=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
